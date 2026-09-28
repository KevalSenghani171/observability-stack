from datetime import datetime, timedelta
from pathlib import Path
import os
import urllib.parse
import urllib.request
import json

from airflow.sdk import DAG, task
from airflow.providers.smtp.hooks.smtp import SmtpHook


# ============================================================
# CONFIGURATION
# ============================================================

# Prometheus URL.
# Change this to your Prometheus service URL.
PROMETHEUS_URL = os.getenv(
    "PROMETHEUS_URL",
    "http://prometheus-server.devops-tools.svc.cluster.local:9090"
)

# SMTP Airflow Connection ID
SMTP_CONN_ID = "smtp_default"

# Email recipients
EMAIL_TO = [
    "senghanikeval@gmail.com"
]

# Report period
REPORT_HOURS = 1


# ============================================================
# PROMETHEUS QUERY FUNCTION
# ============================================================

def prometheus_query(query, start, end, step="300s"):
    """
    Execute Prometheus range query.
    """

    params = urllib.parse.urlencode(
        {
            "query": query,
            "start": start,
            "end": end,
            "step": step,
        }
    )

    url = f"{PROMETHEUS_URL.rstrip('/')}/api/v1/query_range?{params}"

    with urllib.request.urlopen(url, timeout=120) as response:
        data = json.loads(response.read().decode("utf-8"))

    if data.get("status") != "success":
        raise RuntimeError(
            f"Prometheus query failed: {data}"
        )

    return data["data"]["result"]


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def calculate_stats(values):
    """
    Return min/max/avg.
    """

    values = [float(v) for v in values if v is not None]

    if not values:
        return {
            "min": None,
            "max": None,
            "avg": None,
        }

    return {
        "min": min(values),
        "max": max(values),
        "avg": sum(values) / len(values),
    }


def write_excel(report_file, report):
    """
    Create Excel workbook without requiring pandas.
    """

    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError(
            "openpyxl is not installed. "
            "Install apache-airflow with openpyxl in the Airflow image."
        )

    wb = Workbook()

    # --------------------------------------------------------
    # Summary
    # --------------------------------------------------------

    ws = wb.active
    ws.title = "Summary"

    headers = [
        "Server",
        "Metric",
        "Min",
        "Max",
        "Average",
        "Unit",
    ]

    ws.append(headers)

    for row in report["summary"]:
        ws.append(row)

    # --------------------------------------------------------
    # CPU
    # --------------------------------------------------------

    ws = wb.create_sheet("CPU")

    ws.append(
        [
            "Server",
            "CPU Min %",
            "CPU Max %",
            "CPU Average %",
        ]
    )

    for row in report["cpu"]:
        ws.append(row)

    # --------------------------------------------------------
    # Memory
    # --------------------------------------------------------

    ws = wb.create_sheet("Memory")

    ws.append(
        [
            "Server",
            "Memory Min %",
            "Memory Max %",
            "Memory Average %",
        ]
    )

    for row in report["memory"]:
        ws.append(row)

    # --------------------------------------------------------
    # Filesystem
    # --------------------------------------------------------

    ws = wb.create_sheet("Filesystem")

    ws.append(
        [
            "Server",
            "Mountpoint",
            "Min %",
            "Max %",
            "Average %",
        ]
    )

    for row in report["filesystem"]:
        ws.append(row)

    # --------------------------------------------------------
    # Uptime
    # --------------------------------------------------------

    ws = wb.create_sheet("Uptime")

    ws.append(
        [
            "Server",
            "Uptime Min Hours",
            "Uptime Max Hours",
            "Uptime Average Hours",
        ]
    )

    for row in report["uptime"]:
        ws.append(row)

    # --------------------------------------------------------
    # Formatting
    # --------------------------------------------------------

    for ws in wb.worksheets:

        ws.freeze_panes = "A2"

        for cell in ws[1]:
            cell.font = cell.font.copy(bold=True)

        for column in ws.columns:

            max_length = 0

            column_letter = get_column_letter(
                column[0].column
            )

            for cell in column:

                if cell.value is not None:
                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            ws.column_dimensions[
                column_letter
            ].width = min(max_length + 2, 50)

    wb.save(report_file)


# ============================================================
# DAG
# ============================================================

with DAG(
    dag_id="linux_server_utilization_report",

    description=(
        "Generate Linux server CPU, memory, "
        "filesystem and uptime utilization report every 5 minutes "
        "from Prometheus/node_exporter"
    ),

    start_date=datetime(2026, 9, 27),

    schedule="*/5 * * * *",

    catchup=False,

    tags=[
        "linux",
        "prometheus",
        "node-exporter",
        "report",
    ],

    default_args={
        "owner": "devops",
        "retries": 2,
        "retry_delay": timedelta(minutes=5),
    },

) as dag:

    @task
    def generate_report():

        # ----------------------------------------------------
        # Report time range
        # ----------------------------------------------------

        end = datetime.utcnow()

        start = end - timedelta(
            hours=REPORT_HOURS
        )

        start_ts = start.timestamp()
        end_ts = end.timestamp()

        report = {
            "summary": [],
            "cpu": [],
            "memory": [],
            "filesystem": [],
            "uptime": [],
        }

        # ====================================================
        # CPU UTILIZATION
        # ====================================================

        cpu_query = """
        100 -
        (
          avg by (instance)
          (
            rate(node_cpu_seconds_total{
              mode="idle"
            }[5m])
          ) * 100
        )
        """

        cpu_results = prometheus_query(
            cpu_query,
            start_ts,
            end_ts,
        )

        for result in cpu_results:

            instance = result["metric"].get(
                "instance",
                "unknown"
            )

            values = [
                float(v[1])
                for v in result["values"]
            ]

            stats = calculate_stats(values)

            report["cpu"].append(
                [
                    instance,
                    round(stats["min"], 2),
                    round(stats["max"], 2),
                    round(stats["avg"], 2),
                ]
            )

            report["summary"].append(
                [
                    instance,
                    "CPU",
                    round(stats["min"], 2),
                    round(stats["max"], 2),
                    round(stats["avg"], 2),
                    "%",
                ]
            )

        # ====================================================
        # MEMORY UTILIZATION
        # ====================================================

        memory_query = """
        (
          1 -
          (
            node_memory_MemAvailable_bytes
            /
            node_memory_MemTotal_bytes
          )
        ) * 100
        """

        memory_results = prometheus_query(
            memory_query,
            start_ts,
            end_ts,
        )

        for result in memory_results:

            instance = result["metric"].get(
                "instance",
                "unknown"
            )

            values = [
                float(v[1])
                for v in result["values"]
            ]

            stats = calculate_stats(values)

            report["memory"].append(
                [
                    instance,
                    round(stats["min"], 2),
                    round(stats["max"], 2),
                    round(stats["avg"], 2),
                ]
            )

            report["summary"].append(
                [
                    instance,
                    "Memory",
                    round(stats["min"], 2),
                    round(stats["max"], 2),
                    round(stats["avg"], 2),
                    "%",
                ]
            )

        # ====================================================
        # FILESYSTEM / MOUNTPOINT UTILIZATION
        # ====================================================

        filesystem_query = """
        (
          1 -
          (
            node_filesystem_avail_bytes{
              fstype!~"tmpfs|overlay|squashfs"
            }
            /
            node_filesystem_size_bytes{
              fstype!~"tmpfs|overlay|squashfs"
            }
          )
        ) * 100
        """

        filesystem_results = prometheus_query(
            filesystem_query,
            start_ts,
            end_ts,
        )

        for result in filesystem_results:

            metric = result["metric"]

            instance = metric.get(
                "instance",
                "unknown"
            )

            mountpoint = metric.get(
                "mountpoint",
                "unknown"
            )

            values = [
                float(v[1])
                for v in result["values"]
            ]

            stats = calculate_stats(values)

            report["filesystem"].append(
                [
                    instance,
                    mountpoint,
                    round(stats["min"], 2),
                    round(stats["max"], 2),
                    round(stats["avg"], 2),
                ]
            )

            report["summary"].append(
                [
                    instance,
                    f"Filesystem {mountpoint}",
                    round(stats["min"], 2),
                    round(stats["max"], 2),
                    round(stats["avg"], 2),
                    "%",
                ]
            )

        # ====================================================
        # UPTIME
        # ====================================================

        uptime_query = """
        node_time_seconds - node_boot_time_seconds
        """

        uptime_results = prometheus_query(
            uptime_query,
            start_ts,
            end_ts,
        )

        for result in uptime_results:

            instance = result["metric"].get(
                "instance",
                "unknown"
            )

            values = [
                float(v[1]) / 3600
                for v in result["values"]
            ]

            stats = calculate_stats(values)

            report["uptime"].append(
                [
                    instance,
                    round(stats["min"], 2),
                    round(stats["max"], 2),
                    round(stats["avg"], 2),
                ]
            )

            report["summary"].append(
                [
                    instance,
                    "Uptime",
                    round(stats["min"], 2),
                    round(stats["max"], 2),
                    round(stats["avg"], 2),
                    "hours",
                ]
            )

        # ====================================================
        # CREATE EXCEL FILE
        # ====================================================

        report_directory = Path(
            "/tmp/airflow-reports"
        )

        report_directory.mkdir(
            parents=True,
            exist_ok=True
        )

        report_file = (
            report_directory
            / f"linux_server_utilization_"
              f"{end.strftime('%Y%m%d_%H%M%S')}.xlsx"
        )

        write_excel(
            str(report_file),
            report
        )

        if not report_file.exists():
            raise FileNotFoundError(
                f"Excel report was not created: {report_file}"
            )

        print(f"Report generated: {report_file}")

        # Send the report from this same task. This is important when
        # Airflow tasks run in separate Kubernetes pods: /tmp is local
        # to each pod and is not shared between tasks.
        smtp_hook = SmtpHook(smtp_conn_id=SMTP_CONN_ID)

        with smtp_hook:
            smtp_hook.send_email_smtp(
                to=EMAIL_TO,
                subject="Linux Server Utilization Report",
                html_content="""
                    <h3>Linux Server Utilization Report</h3>
                    <p>Please find the latest server utilization report attached.</p>
                """,
                files=[str(report_file)],
            )

        print(f"Report emailed to: {', '.join(EMAIL_TO)}")


